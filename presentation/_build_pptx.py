"""Build presentation/Receipt-to-Cashback.pptx from the 7-slide outline.

Maps 1:1 to the bootcamp brief's 5 sections:
  1. Project Overview      -> Slides 1 (title) + 2 (overview)
  2. Tools & Technologies  -> Slide 4
  3. Solution & Benefits   -> Slides 3 (pipeline) + 5 (numbers)
  4. Challenges Faced      -> Slide 6
  5. Future Steps          -> Slide 7

Ethics is NOT a slide — it lives in docs/04_ethics.md and the in-app
Ethics tab (which is what the project brief actually asks for).

The output is intentionally minimalistic: text + simple tables, no
images. Upload to Google Slides (File -> Open -> Upload) and apply
the bootcamp template colours/visuals from there. Faster than
fighting matplotlib charts inside python-pptx.
"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
OUTPUT_PATH = ROOT / "presentation" / "Receipt-to-Cashback.pptx"

# 16:9 standard slide
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# Brand-ish palette (override in Google Slides if you want)
COLOR_PRIMARY = RGBColor(0x1E, 0x40, 0xAF)   # blue
COLOR_ACCENT = RGBColor(0x10, 0xB9, 0x81)    # green
COLOR_TEXT = RGBColor(0x1F, 0x29, 0x37)
COLOR_MUTED = RGBColor(0x6B, 0x72, 0x80)
COLOR_BG = RGBColor(0xFF, 0xFF, 0xFF)


def add_blank_slide(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])  # Blank


def add_textbox(slide, *, left, top, width, height,
                text, size=18, bold=False,
                color=COLOR_TEXT, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(left, top, width, height)
    tf = tb.text_frame
    tf.word_wrap = True
    # First line goes into the existing default paragraph
    lines = text.split("\n")
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        run = para.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return tb


def add_title(slide, text, *, color=COLOR_PRIMARY):
    add_textbox(
        slide,
        left=Inches(0.6), top=Inches(0.4),
        width=Inches(12), height=Inches(1),
        text=text, size=32, bold=True, color=color,
    )


def add_subtitle(slide, text):
    add_textbox(
        slide,
        left=Inches(0.6), top=Inches(1.2),
        width=Inches(12), height=Inches(0.6),
        text=text, size=16, color=COLOR_MUTED,
    )


def add_accent_bar(slide, *, top=Inches(1.15), color=COLOR_ACCENT):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), top,
                                 Inches(1.6), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = color
    bar.line.fill.background()


def add_two_column_text(slide, *, left_title, left_body, right_title, right_body, top=Inches(2.0)):
    col_w = Inches(5.9)
    # left column
    add_textbox(slide, left=Inches(0.6), top=top, width=col_w, height=Inches(0.6),
                text=left_title, size=18, bold=True, color=COLOR_PRIMARY)
    add_textbox(slide, left=Inches(0.6), top=top + Inches(0.7),
                width=col_w, height=Inches(4.5),
                text=left_body, size=15, color=COLOR_TEXT)
    # right column
    add_textbox(slide, left=Inches(6.85), top=top, width=col_w, height=Inches(0.6),
                text=right_title, size=18, bold=True, color=COLOR_PRIMARY)
    add_textbox(slide, left=Inches(6.85), top=top + Inches(0.7),
                width=col_w, height=Inches(4.5),
                text=right_body, size=15, color=COLOR_TEXT)


def add_table(slide, *, left, top, width, height, data, header_color=COLOR_PRIMARY,
              header_text_color=RGBColor(0xFF, 0xFF, 0xFF), body_font_size=12):
    rows, cols = len(data), len(data[0])
    table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
    table = table_shape.table

    for r_idx, row in enumerate(data):
        for c_idx, cell_text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            tf = cell.text_frame
            tf.word_wrap = True
            para = tf.paragraphs[0]
            run = para.add_run()
            run.text = str(cell_text)
            run.font.size = Pt(body_font_size if r_idx > 0 else body_font_size + 1)
            run.font.bold = (r_idx == 0)
            run.font.color.rgb = header_text_color if r_idx == 0 else COLOR_TEXT
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_color if r_idx == 0 else COLOR_BG
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
    return table


# ---------------------------------------------------------------------------


prs = Presentation()
prs.slide_width = SLIDE_W
prs.slide_height = SLIDE_H


# Slide 1 — Title ------------------------------------------------------------
s = add_blank_slide(prs)
add_textbox(
    s, left=Inches(0.6), top=Inches(2.4), width=Inches(12), height=Inches(1.4),
    text="Receipt-to-Cashback", size=54, bold=True, color=COLOR_PRIMARY,
)
add_textbox(
    s, left=Inches(0.6), top=Inches(3.5), width=Inches(12), height=Inches(0.8),
    text="Scan a receipt. Get paid for the data.",
    size=24, color=COLOR_ACCENT,
)
add_textbox(
    s, left=Inches(0.6), top=Inches(5.6), width=Inches(12), height=Inches(0.6),
    text="Alex Goldbaum", size=18, bold=True,
)
add_textbox(
    s, left=Inches(0.6), top=Inches(6.05), width=Inches(12), height=Inches(0.6),
    text="GenAI & ML Bootcamp 2026 · Developers Institute · Demo Day 11/06/2026",
    size=14, color=COLOR_MUTED,
)
add_textbox(
    s, left=Inches(0.6), top=Inches(6.55), width=Inches(12), height=Inches(0.5),
    text="github.com/alexgl77/receipt-to-cashback   |   huggingface.co/spaces/alexgl77/receipt-to-cashback",
    size=12, color=COLOR_MUTED,
)


# Slide 2 — Project Overview -------------------------------------------------
s = add_blank_slide(prs)
add_title(s, "Project Overview")
add_subtitle(s, "What we built, why, for whom, and what problem it solves")
add_accent_bar(s)

add_two_column_text(
    s,
    left_title="The problem",
    left_body=(
        "Market research today buys self-reported surveys, or shopper "
        "panels that take months to recruit.\n\n"
        "Loyalty programs see one retailer at a time.\n\n"
        "Nobody has the cross-merchant, item-level picture of what "
        "consumers actually buy."
    ),
    right_title="What we built",
    right_body=(
        "A market-research data-acquisition app.\n\n"
        "Users scan a receipt; we pay them a flat cashback in exchange "
        "for the itemised consumption data, which is anonymised and "
        "sold in aggregate.\n\n"
        "Target users:\n"
        "  - consumers who want any-merchant cashback\n"
        "  - brands, retailers and market-research firms (data buyers)"
    ),
)


# Slide 3 — Solution & Pipeline ---------------------------------------------
s = add_blank_slide(prs)
add_title(s, "The solution — pipeline in five stages")
add_subtitle(s, "All pre-trained models stitched with a small amount of business logic")
add_accent_bar(s)

pipeline_lines = (
    "1.  EasyOCR              -> raw text from the receipt photo\n"
    "2.  Gemini 2.5 Flash Lite -> structured JSON (items, prices, total)\n"
    "                            with few-shot prompts + Pydantic schema\n"
    "3.  FAISS + multilingual  -> match each item against the 110-SKU catalog\n"
    "    sentence-transformer    classify spend by category\n"
    "4.  CashbackEngine        -> apply strategy (Flat / Tiered)\n"
    "                            + two-tier drift guard\n"
    "5.  Streamlit UI          -> upload page, B2B analytics, Ethics tab"
)
add_textbox(
    s, left=Inches(0.6), top=Inches(2.0), width=Inches(12.0), height=Inches(4.5),
    text=pipeline_lines, size=18, font="Consolas",
)

add_textbox(
    s, left=Inches(0.6), top=Inches(6.5), width=Inches(12), height=Inches(0.5),
    text="Total end-to-end latency on a real receipt: ~8 seconds (CPU).",
    size=14, color=COLOR_MUTED,
)


# Slide 4 — Tools & Technologies --------------------------------------------
s = add_blank_slide(prs)
add_title(s, "Tools & Technologies")
add_subtitle(s, "What we used — and where in the bootcamp it comes from")
add_accent_bar(s)

tools_table = [
    ["Layer", "Tool", "Source"],
    ["OCR", "EasyOCR (CRAFT detector + CRNN)", "Week 7 — LLM & Gen AI"],
    ["LLM extraction", "Gemini 2.5 Flash Lite + Pydantic schema", "Week 9 — Prompt Engineering"],
    ["Vector matching", "FAISS + multilingual sentence-transformer", "Week 8 — NLP & RAG"],
    ["Clustering", "scikit-learn KMeans + PCA (visual)", "Week 4–5 — ML & Stats"],
    ["A/B stats", "scipy Welch's t-test", "Week 5 — Statistics for ML"],
    ["Data wrangling", "pandas, matplotlib, seaborn", "Week 3–4 — Data Analysis"],
    ["Code style", "Python OOP, dataclasses, Pydantic", "Week 1–2 — Python & OOP"],
    ["Web app", "Streamlit (4 pages)", "SELF-TAUGHT (gap from bootcamp)"],
]
add_table(
    s,
    left=Inches(0.6), top=Inches(2.0),
    width=Inches(12.0), height=Inches(4.8),
    data=tools_table,
    body_font_size=14,
)


# Slide 5 — Solution and Benefits (what works, in numbers) ------------------
s = add_blank_slide(prs)
add_title(s, "What works — by the numbers")
add_subtitle(s, "Measured on the CORD-v2 dataset (1000 real receipts, Naver Clova 2019)")
add_accent_bar(s)

metrics_table = [
    ["Metric", "Value", "Note"],
    ["End-to-end latency",      "~8 s / receipt",    "CPU, no GPU"],
    ["OCR loose recall",         "85% (100% on 16/20)", "20-receipt benchmark"],
    ["Catalog classification",   "100% via multilingual embedding", "EN + ID + KR mix"],
    ["Unit tests",               "13 / 13 passing",   "CashbackEngine"],
    ["Drift guard tiers",        "10% scale · 50% refuse", "prevents over- and under-pay"],
    ["A/B test result (synthetic users)", "+29% receipts at 3% vs 2% cashback", "p < 0.05, Welch's t-test"],
    ["K-Means clustering",       "4 archetypes recovered cleanly", "café · family · sweet · office"],
    ["Live demo",                "huggingface.co/spaces/alexgl77/receipt-to-cashback", ""],
]
add_table(
    s,
    left=Inches(0.6), top=Inches(2.0),
    width=Inches(12.0), height=Inches(4.8),
    data=metrics_table,
    body_font_size=14,
)
add_textbox(
    s, left=Inches(0.6), top=Inches(7.0), width=Inches(12), height=Inches(0.4),
    text="Live demo next (video).",
    size=14, color=COLOR_ACCENT, bold=True,
)


# Slide 6 — Challenges Faced -------------------------------------------------
s = add_blank_slide(prs)
add_title(s, "Challenges faced — and how we fixed them")
add_subtitle(s, "Four problems found mid-build, not at the demo")
add_accent_bar(s)

challenges_table = [
    ["#", "Problem", "Fix"],
    ["1", "Original business model (partner-rebate) rejected legitimate items",
          "Pivoted to market-research: every priced line earns cashback; catalog became a classifier"],
    ["2", "English-only embedding rejected an Indonesian item (PKT AYAM)",
          "Swapped to multilingual sentence-transformer (paraphrase-multilingual-MiniLM-L12-v2)"],
    ["3", "End-to-end latency was 40+ s on the first run",
          "Switched to Gemini Flash Lite + cached OCR/LLM by image hash → 8 s"],
    ["4", "Rehearsal: LLM double-counted items, then mis-read 1,591,600 as 591,600 (inverse failure)",
          "Two-tier drift guard: scale at 10% drift, refuse cashback at 50% drift"],
]
add_table(
    s,
    left=Inches(0.6), top=Inches(2.0),
    width=Inches(12.0), height=Inches(5.0),
    data=challenges_table,
    body_font_size=13,
)


# Slide 7 — Future Steps + Thanks -------------------------------------------
s = add_blank_slide(prs)
add_title(s, "Future steps & thanks")
add_subtitle(s, "What we would add with two more weeks — and who helped us get here")
add_accent_bar(s)

add_two_column_text(
    s,
    left_title="If we had two more weeks",
    left_body=(
        "·  Real users (not synthetic)\n"
        "·  Postgres warehouse for receipt records\n"
        "·  Image-level redaction (cardholder + identifiers)\n"
        "·  K-anonymity floor at sale time\n"
        "·  MCP agent: ask \"how much did I spend on coffee in May?\"\n"
        "·  CNN quality gate (blurry/junk filter)\n"
        "·  Multilingual OCR (Hebrew, Arabic) for global users\n"
        "\nFull ethics discussion (privacy, consent, OCR bias, "
        "hallucination policy, k-anonymity) lives in the repo "
        "(docs/04_ethics.md) and as the Ethics tab inside the app."
    ),
    right_title="Thanks",
    right_body=(
        "Yossi Eikelman — instructor at Developers Institute,\n"
        "whose mid-build feedback \"spine first, garnish after\"\n"
        "set the priority order that made delivery on time possible.\n\n"
        "DI cohort 2026 — sounding board.\n\n"
        "Naver Clova — for releasing CORD-v2 under CC BY 4.0.\n\n"
        "Google AI Studio — Gemini API free tier.\n\n"
        "Hugging Face — multilingual sentence-transformer + Space hosting."
    ),
)

add_textbox(
    s, left=Inches(0.6), top=Inches(7.0), width=Inches(12), height=Inches(0.4),
    text="Code: github.com/alexgl77/receipt-to-cashback   |   Live: huggingface.co/spaces/alexgl77/receipt-to-cashback",
    size=12, color=COLOR_MUTED,
)


# Save
prs.save(OUTPUT_PATH)
print(f"Wrote {OUTPUT_PATH}")
print(f"Slides: {len(prs.slides)}")
